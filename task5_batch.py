"""
Task 5 — Batch Processing & Excel Report
"""

import os
import sys
from datetime import datetime
from pathlib import Path

import cv2
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

sys.path.insert(0, os.path.dirname(__file__))
from task1_qr_decoder import decode_answer_key
from task2_ocr import extract_student_info
from task3_bubble_reader import read_bubble_sheet
from task4_grader import grade_quiz


GRADE_COLORS = {
    "A+": "2ECC71", "A": "27AE60",
    "B":  "3498DB", "C": "F39C12",
    "D":  "E67E22", "F": "E74C3C",
}


def _color_for_grade(grade: str) -> str:
    return GRADE_COLORS.get(grade, "FFFFFF")


def run_batch(
    image_dir: str,
    quiz_label: str = "Quiz 1",
    output_dir: str = "output",
) -> str:
    image_dir = Path(image_dir)
    output_dir = Path(output_dir)
    output_dir.mkdir(exist_ok=True)

    extensions = {".jpg", ".jpeg", ".png", ".bmp", ".tiff"}
    image_paths = sorted([
        p for p in image_dir.iterdir()
        if p.suffix.lower() in extensions
    ])

    if not image_paths:
        raise ValueError(f"No images found in {image_dir}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = [
        "Quiz", "Set", "Class", "Subject", "Name", "Reg No",
        *[f"Part1_Q{i:02d}" for i in range(1, 9)],
        *[f"Part2_Q{i:02d}" for i in range(1, 9)],
        "Correct", "Incorrect", "Unattempted",
        "Total Marks", "Max Marks", "Percentage", "Grade",
    ]

    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(color="FFFFFF", bold=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    answer_key = None
    rows = []
    total_marks_list = []

    for img_path in image_paths:
        try:
            image = cv2.imread(str(img_path))
            if image is None:
                continue

            if answer_key is None:
                answer_key = decode_answer_key(image)

            student_info = extract_student_info(image)
            student_answers = read_bubble_sheet(image)

            if answer_key:
                report = grade_quiz(
                    student_answers, answer_key,
                    student_name=student_info.name,
                    reg_no=student_info.reg_no,
                )
                total_marks_list.append(report.total_marks)
                row = [
                    quiz_label,
                    answer_key.quiz_set,
                    "BSE-4A",
                    "Artificial Intelligence",
                    student_info.name,
                    student_info.reg_no,
                    *[student_answers.part1.get(f"Q{i:02d}", "") or ""
                      for i in range(1, 9)],
                    *[student_answers.part2.get(f"Q{i:02d}", "") or ""
                      for i in range(1, 9)],
                    report.correct,
                    report.incorrect,
                    report.unattempted,
                    report.total_marks,
                    report.max_marks,
                    f"{report.percentage:.1f}%",
                    report.grade,
                ]
            else:
                total_marks_list.append(0)
                row = [
                    quiz_label, "?",
                    "BSE-4A", "Artificial Intelligence",
                    student_info.name, student_info.reg_no,
                    *["" for _ in range(16)],
                    0, 0, 0, 0, 0, "0%", "F",
                ]

            rows.append(row)

        except Exception as e:
            rows.append([
                quiz_label, "ERROR", "BSE-4A", "AI",
                str(img_path.name), str(e),
                *["" for _ in range(19)]
            ])

    alt_fill = PatternFill("solid", fgColor="ECF0F1")

    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(horizontal="center")
            if r_idx % 2 == 0:
                cell.fill = alt_fill

        grade = row[-1] if row else "F"
        grade_fill = PatternFill(
            "solid", fgColor=_color_for_grade(grade))
        grade_font = Font(color="FFFFFF", bold=True)
        grade_cell = ws.cell(row=r_idx, column=len(headers))
        grade_cell.fill = grade_fill
        grade_cell.font = grade_font

    # ✅ Summary row
    summary_row = len(rows) + 2
    summary_fill = PatternFill("solid", fgColor="2C3E50")
    summary_font = Font(color="FFFFFF", bold=True)

    if total_marks_list:
        avg = sum(total_marks_list) / len(total_marks_list)
        highest = max(total_marks_list)
        lowest = min(total_marks_list)
        summary_text = f"Avg:{avg:.1f} High:{highest} Low:{lowest}"
    else:
        summary_text = "No data"

    for c_idx in range(1, len(headers) + 1):
        val = "SUMMARY" if c_idx == 1 else (
            summary_text if c_idx == len(headers) - 3 else "")
        cell = ws.cell(row=summary_row, column=c_idx, value=val)
        cell.fill = summary_fill
        cell.font = summary_font
        cell.alignment = Alignment(horizontal="center")

    # Column widths
    for col in ws.columns:
        max_len = max(
            (len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[
            col[0].column_letter].width = min(max_len + 4, 30)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = output_dir / f"quiz_results_{timestamp}.xlsx"
    wb.save(str(out_path))

    print(f"✓ Saved: {out_path}")
    return str(out_path)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("folder", help="Folder with quiz images")
    parser.add_argument("--quiz", default="Quiz 1")
    parser.add_argument("--output", default="output")
    args = parser.parse_args()
    run_batch(args.folder, quiz_label=args.quiz,
              output_dir=args.output)